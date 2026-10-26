from __future__ import annotations

import hashlib
import json
import math
import re
import subprocess
import xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
NAME = "206题原文与作答审计-2026-10-24"
PATH = ROOT / (NAME + ".xlsx")
CALCULATED = ROOT / (NAME + "_recalc_verify.xlsx") / (NAME + ".XLSX")
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(8 * 1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def close(actual, expected):
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        assert isinstance(actual, (int, float)) and math.isclose(actual, expected, rel_tol=0, abs_tol=1e-10), (actual, expected)
    else:
        assert actual == expected, (actual, expected)


def main():
    rows = [json.loads(line) for line in (ROOT / "joined-206-2026-10-24.jsonl").read_text(encoding="utf-8").splitlines()]
    summary = json.loads((ROOT / "analysis-summary-2026-10-24.json").read_text(encoding="utf-8"))
    wb = load_workbook(PATH, data_only=False)
    calculated = load_workbook(CALCULATED, data_only=True)
    formulas = {}
    for ws in wb.worksheets:
        assert not ws.sheet_view.showGridLines
        for cells in ws.iter_rows():
            for cell in cells:
                if cell.data_type == "f":
                    value = calculated[ws.title][cell.coordinate].value
                    assert value is not None, (ws.title, cell.coordinate, "missing calculation result")
                    assert not isinstance(value, str) or not value.startswith(("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!")), (ws.title, cell.coordinate, value)
                    formulas[(ws.title, cell.coordinate)] = (cell.value, value)
    assert formulas
    class_names = summary["classification"]["confusion_matrix"]["classes"]
    for i, row in enumerate(rows, 7):
        for col, expected in ((3, row["id"]), (14, row["old_text"]), (15, row["new_text"]), (16, row["background"])):
            close(wb["逐题原表"].cell(i, col).value, expected)
        for col, expected in ((9, int(row["correct"])), (10, row["confidence"]), (27, int(row["swap_consistent"])), (34, row["margin"])):
            close(calculated["逐题原表"].cell(i, col).value, expected)
        for col, expected in enumerate(row["probabilities"], 20):
            close(wb["逐题原表"].cell(i, col).value, expected)
    for i, name in enumerate(class_names, 7):
        item = summary["classification"]["per_class"][name]
        for col, expected in ((3, item["support"]), (4, summary["prediction_counts"].get(name, 0)), (5, item["tp"]), (6, item["precision"]), (7, item["recall"]), (8, item["f1"])):
            close(calculated["类别与混淆"].cell(i, col).value, expected)
    for i, line in enumerate(summary["classification"]["confusion_matrix"]["matrix"], 17):
        for col, expected in enumerate(line, 3):
            close(calculated["类别与混淆"].cell(i, col).value, expected)
    cover = {"C7": 206, "C8": 46, "C9": summary["classification"]["accuracy"], "C10": summary["classification"]["macro_f1"],
             "C11": summary["majority_baseline"]["classification"]["accuracy"], "C12": summary["majority_baseline"]["classification"]["macro_f1"],
             "C13": 68, "C14": 0, "C15": 48}
    for address, expected in cover.items():
        close(calculated["封面"][address].value, expected)
    close(calculated["置信度与接管"]["D24"].value, 0)
    close(calculated["置信度与接管"]["F24"].value, "无样本")
    fulltext = defaultdict(list)
    for values in wb["全文与可见输入"].iter_rows(min_row=7, min_col=2, max_col=7, values_only=True):
        identifier, field, part, total, chars, content = values
        assert len(content or "") == chars
        fulltext[(identifier, field)].append((part, total, content or ""))
    for row in rows:
        expected_fields = {"共同背景": row["background"], "旧条目完整JSON": json.dumps(row["pair"]["old"], ensure_ascii=False, indent=2),
                           "新条目完整JSON": json.dumps(row["pair"]["new"], ensure_ascii=False, indent=2),
                           "完整pair JSON": json.dumps(row["pair"], ensure_ascii=False, indent=2),
                           "学生实际可见_正向": row["model_visible_forward"], "学生实际可见_反向": row["model_visible_reverse"]}
        for field, expected in expected_fields.items():
            parts = sorted(fulltext[(row["id"], field)])
            assert len(parts) == parts[0][1]
            assert "".join(item[2] for item in parts) == expected
    html_path = ROOT / "206题可视化与逐题阅读器-2026-10-24.html"
    html = html_path.read_text(encoding="utf-8")
    match = re.search(r'<script id="data" type="application/json">(.*?)</script>', html, re.S)
    payload = json.loads(match.group(1))
    assert len(payload["rows"]) == 206
    for actual, source in zip(payload["rows"], rows):
        for key in ("id", "pair", "gold", "prediction", "probabilities", "model_visible_forward"):
            assert actual[key] == source[key]
    code = re.findall(r"<script>(.*?)</script>", html, re.S)[0]
    node = subprocess.run(["node", "--check", "--input-type=commonjs"], input=code, text=True, capture_output=True, encoding="utf-8")
    assert node.returncode == 0, node.stderr
    original_formulas = dict(formulas)
    with ZipFile(PATH) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    charts = {name: hashlib.sha256(value).hexdigest() for name, value in members.items() if name.startswith("xl/charts/")}
    workbook = ET.fromstring(members["xl/workbook.xml"])
    rels = ET.fromstring(members["xl/_rels/workbook.xml.rels"])
    targets = {item.attrib["Id"]: item.attrib["Target"] for item in rels}
    injected = 0
    for item in workbook.find("s:sheets", NS):
        sheet_name = item.attrib["name"]
        target = targets[item.attrib["{" + NS["r"] + "}id"]]
        member = target.lstrip("/") if target.startswith("/") else "xl/" + target
        xml = ET.fromstring(members[member])
        for cell in xml.findall(".//s:sheetData/s:row/s:c", NS):
            key = (sheet_name, cell.attrib["r"])
            if key not in formulas:
                continue
            expression, value = formulas[key]
            assert cell.find("s:f", NS).text == expression[1:]
            cached = cell.find("s:v", NS)
            if cached is None:
                cached = ET.SubElement(cell, "{" + NS["s"] + "}v")
            if isinstance(value, bool):
                cell.attrib["t"], cached.text = "b", "1" if value else "0"
            elif isinstance(value, (int, float)):
                assert math.isfinite(value)
                cell.attrib["t"], cached.text = "n", repr(value)
            else:
                cell.attrib["t"], cached.text = "str", value
            injected += 1
        members[member] = ET.tostring(xml, encoding="utf-8", xml_declaration=True)
    temp = ROOT / "formula-cache-injection.tmp.xlsx"
    with ZipFile(temp, "w", compression=8) as archive:
        for name, value in members.items():
            archive.writestr(name, value)
    temp.replace(PATH)
    final_formula = load_workbook(PATH, data_only=False)
    final_data = load_workbook(PATH, data_only=True)
    for (sheet_name, coordinate), (expression, value) in original_formulas.items():
        assert final_formula[sheet_name][coordinate].value == expression
        close(final_data[sheet_name][coordinate].value, value)
    assert sum(len(ws._charts) for ws in final_formula.worksheets) == 12
    with ZipFile(PATH) as archive:
        assert charts == {name: hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith("xl/charts/")}
    source_manifest = json.loads((ROOT / "source-manifest-2026-10-24.json").read_text(encoding="utf-8"))
    repo = ROOT.parents[3] / "agent"
    assert all(sha(repo / item["logical_path"]) == item["sha256"] for item in source_manifest["files"])
    result = {"passed": True, "source_files_unchanged": True, "raw_rows": 206, "fulltext_reassembled_without_truncation": True,
              "formula_engine": "formulas", "engine_copy_reconciled_to_frozen_metrics": True, "formulas_cached_without_replacing_expressions": injected,
              "formula_expression_and_chart_xml_unchanged": True, "embedded_excel_charts": 12,
              "html_data_rows_and_all_selected_fields_equal": True, "html_javascript_syntax": "passed", "cover_values": cover,
              "xlsx_sha256": sha(PATH), "html_sha256": sha(html_path)}
    (ROOT / "delivery-verification-2026-10-24.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
