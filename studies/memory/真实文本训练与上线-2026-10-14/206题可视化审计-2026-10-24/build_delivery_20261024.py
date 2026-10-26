from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent
PATH = ROOT / "206题原文与作答审计-2026-10-24.xlsx"
CLASSES = ("equivalent", "old_entails_new_strict", "new_entails_old_strict", "disjoint", "overlap_incomparable")
ZH = dict(zip(CLASSES, ("等价", "旧严格蕴含新", "新严格蕴含旧", "互斥", "相交但互不包含")))
FAMILY = {"comparative_preferences": "比较偏好", "conjunction_disjunction": "合取与析取", "cross_entry_joint_consistency": "跨条目联合一致性",
          "entity_event_coreference": "实体与事件指代", "exceptions_default_rules": "例外与默认规则", "multi_clause_local_edits": "多子句局部编辑",
          "negation_polarity": "否定与极性", "numeric_intervals_units": "数值区间与单位", "order_causality": "顺序与因果", "quantifier_scope": "量词作用域",
          "taxonomy_granularity": "分类粒度", "temporal_aspect_content": "时间与体态"}
BLUE, GRAY, PALE = "0066CC", "666666", "E6F0FA"
COLORS = ("333333", "0066CC", "4A90D9", "888888", "B3D2F2")


def text(ws, row, col, value):
    cell = ws.cell(row, col, value)
    if isinstance(value, str):
        cell.data_type = "s"
    cell.font = Font(name="Microsoft YaHei", size=10, color=BLUE)
    cell.alignment = Alignment(vertical="center", wrap_text=False)
    return cell


def formula(ws, row, col, value, fmt=None):
    cell = ws.cell(row, col, value)
    cell.font = Font(name="Microsoft YaHei", size=10, color="000000")
    cell.alignment = Alignment(vertical="center", wrap_text=False)
    if fmt:
        cell.number_format = fmt
    return cell


def sheet(wb, name, title, note, headers):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.row_dimensions[2].height = 34
    ws.row_dimensions[4].height = 24
    ws.column_dimensions["A"].width = 3
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=max(8, min(len(headers) + 1, 12)))
    text(ws, 2, 2, title).font = Font(name="Microsoft YaHei", size=19, bold=True, color="000000")
    text(ws, 4, 2, note).font = Font(name="Microsoft YaHei", size=10, color=GRAY)
    for i, label in enumerate(headers, 2):
        cell = text(ws, 6, i, label)
        cell.fill = PatternFill("solid", fgColor="333333")
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        ws.column_dimensions[get_column_letter(i)].width = 19
    ws.row_dimensions[6].height = 28
    ws.freeze_panes = "D7"
    return ws


def finish(wb, ws, last_row, last_col, *, table_name=None):
    if table_name:
        table = Table(displayName=table_name, ref=f"B6:{get_column_letter(last_col)}{last_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    for row in range(7, last_row + 1):
        ws.row_dimensions[row].height = 23
    wb.save(PATH)
    check = load_workbook(PATH)
    actual = check[ws.title]
    assert not actual.sheet_view.showGridLines and actual["B2"].value
    assert actual.max_row >= last_row
    assert len(actual._charts) == len(ws._charts)
    for cells in actual.iter_rows():
        for cell in cells:
            if isinstance(cell.value, str):
                assert len(cell.value.encode("utf-16-le")) // 2 <= 32767
    check.close()


def chart(ws, title, min_col, max_col, header_row, last_row, category_col, anchor, *, percent=False, line=False, stacked=False):
    item = LineChart() if line else BarChart()
    item.title = title
    item.height, item.width = 8.2, 18.2
    item.add_data(Reference(ws, min_col=min_col, max_col=max_col, min_row=header_row, max_row=last_row), titles_from_data=True)
    item.set_categories(Reference(ws, min_col=category_col, min_row=header_row + 1, max_row=last_row))
    if not line:
        item.type = "col"
        item.grouping = "stacked" if stacked else "clustered"
        if stacked:
            item.overlap = 100
    if percent:
        item.y_axis.numFmt = "0%"
        item.y_axis.scaling.min, item.y_axis.scaling.max = 0, 1
    for i, series in enumerate(item.series):
        series.graphicalProperties.solidFill = COLORS[i % len(COLORS)]
        series.graphicalProperties.line.solidFill = COLORS[i % len(COLORS)]
    ws.add_chart(item, anchor)


def main():
    rows = [json.loads(line) for line in (ROOT / "joined-206-2026-10-24.jsonl").read_text(encoding="utf-8").splitlines()]
    summary = json.loads((ROOT / "analysis-summary-2026-10-24.json").read_text(encoding="utf-8"))
    sources = json.loads((ROOT / "source-manifest-2026-10-24.json").read_text(encoding="utf-8"))
    frame = pd.DataFrame(rows)
    assert len(frame) == frame["id"].nunique() == 206 and int(frame["correct"].sum()) == 46
    wb = Workbook()
    wb.remove(wb.active)
    headers = ["序号", "题ID", "分区", "问题族", "条目类型", "AI参考标签", "模型作答", "与参考一致(1/0)", "最大softmax(未校准)", "输入完整性", "原完整输入tokens", "实际可见tokens", "旧条目正文", "新条目正文", "共同背景", "旧条目完整JSON", "新条目完整JSON", "完整pair JSON", *["概率_" + ZH[c] for c in CLASSES], "反向AI参考", "反向模型作答", "交换一致(1/0)", "依赖簇", "语言", "语义模式", "来源", "参考性质", "可见输入复现", "前二概率差", "人工复核标签", "人工复核备注"]
    ws = sheet(wb, "逐题原表", "206题 · 原文与冻结作答", "正向每root一次。AI参考不是人工金标；数值1/0表示与参考一致，不代表已确认事实。末两列供人工复核，不参与原分数。", headers)
    for i, row in enumerate(rows, 7):
        fields = [row["number"], row["id"], row["split"], FAMILY[row["family"]], row["kind"], row["gold_zh"], row["prediction_zh"], None, None,
                  "完整" if row["complete"] else "截断", row["full_tokens"], row["visible_tokens"], row["old_text"], row["new_text"], row["background"],
                  json.dumps(row["pair"]["old"], ensure_ascii=False), json.dumps(row["pair"]["new"], ensure_ascii=False), json.dumps(row["pair"], ensure_ascii=False),
                  *row["probabilities"], ZH[row["reverse_gold"]], ZH[row["reverse_prediction"]], None, row["cluster_id"], row["language"], row["mode"], row["source_id"],
                  "AI代理；非人工金标", "view/encoding SHA均相同", None, "", ""]
        for col, value in enumerate(fields, 2):
            if value is not None:
                text(ws, i, col, value)
        formula(ws, i, 9, f"=IF(G{i}=H{i},1,0)", "0")
        formula(ws, i, 10, f"=MAX(T{i}:X{i})", "0.0%")
        formula(ws, i, 27, f'=IF(Z{i}=IF(H{i}="旧严格蕴含新","新严格蕴含旧",IF(H{i}="新严格蕴含旧","旧严格蕴含新",H{i})),1,0)', "0")
        formula(ws, i, 34, f"=LARGE(T{i}:X{i},1)-LARGE(T{i}:X{i},2)", "0.0000")
        for col in range(20, 25):
            ws.cell(i, col).number_format = "0.0000"
    for col in ("C", "AB"):
        ws.column_dimensions[col].width = 29
    for col in ("N", "O"):
        ws.column_dimensions[col].width = 49
    for col in ("P", "Q", "R", "S"):
        ws.column_dimensions[col].width = 65
    ws.conditional_formatting.add("J7:J212", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="4A90D9"))
    finish(wb, ws, 212, len(headers) + 1, table_name="QuestionRows")
    ranges = {"gold": "'逐题原表'!$G$7:$G$212", "pred": "'逐题原表'!$H$7:$H$212", "correct": "'逐题原表'!$I$7:$I$212", "confidence": "'逐题原表'!$J$7:$J$212", "complete": "'逐题原表'!$K$7:$K$212"}

    ws = sheet(wb, "全文与可见输入", "全文与学生实际可见输入", "每个字段完整保存；超过15000字符时分段。按题ID+字段+段号拼接即原文。内容来自原始题及摘要核同的重建输入，不补造模型理由。", ["题ID", "字段", "段号", "总段数", "字符数", "全文内容"])
    cursor = 7
    for row in rows:
        content = {"共同背景": row["background"], "旧条目完整JSON": json.dumps(row["pair"]["old"], ensure_ascii=False, indent=2),
                   "新条目完整JSON": json.dumps(row["pair"]["new"], ensure_ascii=False, indent=2), "完整pair JSON": json.dumps(row["pair"], ensure_ascii=False, indent=2),
                   "学生实际可见_正向": row["model_visible_forward"], "学生实际可见_反向": row["model_visible_reverse"],
                   "参考标签来源": json.dumps({"label_origin": row["label_origin"], "NL_fidelity_origin": row["nl_fidelity_origin"], "adjudication_origin": row["adjudication_origin"], "teacher_evidence": row["teacher_evidence"], "evidence": row["evidence"]}, ensure_ascii=False, indent=2)}
        for field, value in content.items():
            chunks = [value[j:j + 15000] for j in range(0, len(value), 15000)] or [""]
            for n, chunk in enumerate(chunks, 1):
                for col, item in enumerate((row["id"], field, n, len(chunks), len(chunk), chunk), 2):
                    text(ws, cursor, col, item)
                cursor += 1
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["G"].width = 150
    finish(wb, ws, cursor - 1, 7, table_name="FullTextRows")

    ws = sheet(wb, "类别与混淆", "类别分布 · 混淆矩阵", "行=AI参考，列=模型作答；对角线为与参考一致。支持仅4题的互斥类不能用于风险认证。", ["类别", "参考数量", "预测数量", "一致数量TP", "精确率", "召回率", "F1"])
    for i, name in enumerate(CLASSES, 7):
        text(ws, i, 2, ZH[name])
        formula(ws, i, 3, f'=COUNTIF({ranges["gold"]},B{i})', "0")
        formula(ws, i, 4, f'=COUNTIF({ranges["pred"]},B{i})', "0")
        formula(ws, i, 5, f'=COUNTIFS({ranges["gold"]},B{i},{ranges["pred"]},B{i})', "0")
        formula(ws, i, 6, f"=IFERROR(E{i}/D{i},0)", "0.0%")
        formula(ws, i, 7, f"=IFERROR(E{i}/C{i},0)", "0.0%")
        formula(ws, i, 8, f"=IFERROR(2*F{i}*G{i}/(F{i}+G{i}),0)", "0.0000")
    for i, name in enumerate(CLASSES, 3):
        text(ws, 16, i, ZH[name])
    text(ws, 16, 2, "参考↓ / 预测→")
    for i, name in enumerate(CLASSES, 17):
        text(ws, i, 2, ZH[name])
        for col in range(3, 8):
            formula(ws, i, col, f'=COUNTIFS({ranges["gold"]},$B{i},{ranges["pred"]},{get_column_letter(col)}$16)', "0")
    ws.conditional_formatting.add("C17:G21", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color="4A90D9"))
    chart(ws, "参考分布与模型预测分布", 3, 4, 6, 11, 2, "K3")
    chart(ws, "各类别召回率", 7, 7, 6, 11, 2, "K19", percent=True)
    chart(ws, "每个参考类别被判成什么", 3, 7, 16, 21, 2, "B26", stacked=True)
    ws.column_dimensions["B"].width = 27
    finish(wb, ws, 21, 8)

    ws = sheet(wb, "分层分析", "完整性、分区与题型分层", "所有比例分母均在旁列；分组不是随机分配，不能把差异或相近分数当作因果效应。题号也不是训练时间趋势。", ["维度", "分组", "题数n", "一致数", "准确率", "样本提示"])
    definitions = [("complete", "输入完整性", "K", lambda value: "完整" if value else "截断"), ("split", "评测分区", "D", str),
                   ("kind", "条目类型", "F", str), ("language", "语言", "AC", str), ("mode", "语义模式", "AD", str),
                   ("family", "问题族", "E", lambda value: FAMILY[value])]
    cursor, group_ranges = 7, {}
    for field, axis, letter, show in definitions:
        start = cursor
        for item in summary["groups"][field]:
            text(ws, cursor, 2, axis)
            text(ws, cursor, 3, show(item["value"]))
            ref = f"'逐题原表'!${letter}$7:${letter}$212"
            formula(ws, cursor, 4, f"=COUNTIF({ref},C{cursor})", "0")
            formula(ws, cursor, 5, f'=SUMIF({ref},C{cursor},{ranges["correct"]})', "0")
            formula(ws, cursor, 6, f"=IFERROR(E{cursor}/D{cursor},0)", "0.0%")
            formula(ws, cursor, 7, f'=IF(D{cursor}<10,"小样本；谨慎解读","描述性分层")')
            cursor += 1
        group_ranges[field] = (start, cursor - 1)
    for field, anchor in (("complete", "J3"), ("split", "J19"), ("kind", "J35"), ("language", "B40"), ("mode", "J51"), ("family", "B57")):
        first, last = group_ranges[field]
        header = cursor + 2
        text(ws, header, 2, "分组")
        text(ws, header, 3, "准确率")
        for offset, row in enumerate(range(first, last + 1), header + 1):
            formula(ws, offset, 2, f"=C{row}")
            formula(ws, offset, 3, f"=F{row}", "0.0%")
        chart(ws, definitions[[x[0] for x in definitions].index(field)][1] + "：准确率", 3, 3, header, header + last - first + 1, 2, anchor, percent=True)
        cursor = header + last - first + 3
    finish(wb, ws, cursor - 2, 7)

    ws = sheet(wb, "置信度与接管", "未校准置信度 · 门控覆盖", "max-softmax不是正确概率。下表仅重算旧评测上的阈值描述，不选择新阈值；生产还要求完整输入及其他检查。", ["置信度下界", "上界(不含)", "区间", "一致数", "不一致数", "合计n"])
    bins = [(0, .25), (.25, .30), (.30, .35), (.35, .40), (.40, 1.000001)]
    for i, (lo, hi) in enumerate(bins, 7):
        text(ws, i, 2, lo)
        text(ws, i, 3, hi)
        text(ws, i, 4, f"{lo:.0%}–{min(hi,1):.0%}")
        for col, correct in ((5, 1), (6, 0)):
            formula(ws, i, col, f'=COUNTIFS({ranges["confidence"]},">="&B{i},{ranges["confidence"]},"<"&C{i},{ranges["correct"]},{correct})', "0")
        formula(ws, i, 7, f"=SUM(E{i}:F{i})", "0")
    for col, value in enumerate(("阈值", "只看分数通过n", "完整且分数通过n", "其中一致n", "接受部分准确率", "完整接管覆盖率"), 2):
        text(ws, 17, col, value)
    for i, threshold in enumerate((.20, .25, .30, .35, .40, .50, .80), 18):
        text(ws, i, 2, threshold).number_format = "0%"
        formula(ws, i, 3, f'=COUNTIF({ranges["confidence"]},">="&B{i})', "0")
        formula(ws, i, 4, f'=COUNTIFS({ranges["confidence"]},">="&B{i},{ranges["complete"]},"完整")', "0")
        formula(ws, i, 5, f'=SUMIFS({ranges["correct"]},{ranges["confidence"]},">="&B{i},{ranges["complete"]},"完整")', "0")
        formula(ws, i, 6, f'=IF(D{i}=0,"无样本",E{i}/D{i})', "0.0%")
        formula(ws, i, 7, f'=D{i}/COUNTA(\'逐题原表\'!$C$7:$C$212)', "0.0%")
    chart(ws, "正确与错误的置信度分布", 5, 6, 6, 11, 4, "J3", stacked=True)
    chart(ws, "完整输入的描述性接管覆盖率", 7, 7, 17, 24, 2, "J20", line=True, percent=True)
    finish(wb, ws, 24, 7)

    ws = sheet(wb, "历史模型对照", "同206题的历史模型结果", "存档12模型描述结果，不是本轮新训练；不据本表重选赢家。dual3非法组合仍计错；不能直接跨头比较其条件概率分数。", ["模型", "输入版本", "头", "seed", "一致数(存档)", "题数n", "准确率", "宏F1(存档)", "完整题数", "完整子集准确率(存档)"])
    for i, model in enumerate(summary["archived_models"], 7):
        values = (model["model_id"], model["version"], model["arm"], model["seed"], model["correct"], 206, None, model["macro_f1"], model["input_complete_roots"], model["truncation"]["complete"]["accuracy"])
        for col, value in enumerate(values, 2):
            if value is not None:
                text(ws, i, col, value)
        formula(ws, i, 8, f"=F{i}/G{i}", "0.0%")
        ws.cell(i, 9).number_format = "0.0000"
        ws.cell(i, 11).number_format = "0.0%"
    chart(ws, "同题历史准确率与宏F1", 8, 9, 6, 18, 2, "M3", percent=True)
    ws.column_dimensions["B"].width = 28
    finish(wb, ws, 18, 11, table_name="ArchiveModels")

    ws = sheet(wb, "来源与标签说明", "口径、来源与字段解释", "全部来源为本机冻结研究工件，无外部网络采集；这里只列相对逻辑路径与SHA，不包含凭据或真实运维信息。", ["来源名称 / 字段", "位置 / 解释", "SHA256", "字节数"])
    docs = [("模型", "v3-c5-seed13的冻结历史作答，非今天重新运行的线上响应；旧评测导出包与新runtime封装hash不同，不等于参数不同。"),
            ("标答性质", "206条均为blind_teacher_proxy，无人工金标；pending_independent_adjudication不代表已经人工裁决。"),
            ("old / new方向", "old=P，new=H。旧严格蕴含新表示旧约束更强，新严格蕴含旧表示新约束更强；不自动决定覆盖权限。"),
            ("样本依赖", "206个root属于43个依赖簇。反向和同簇题不是新的独立样本。"),
            ("多数类基线", "事后始终猜overlap_incomparable，仅作标签分布参照；其余四类召回为零。"),
            ("hard_success", "原字段仅表示生成合法分类，不表示答对；本表用与AI参考标签一致单独计算。"),
            ("不生成理由", "BERT输出类别/概率，不生成自然语言理由；本交付没有补造模型思维链。"),
            ("复核入口", "建议先用离线HTML阅读器看全文/可见输入，再在逐题表末两列填写自己的复核意见；不覆盖原标签。")]
    for i, (name, content) in enumerate(docs, 7):
        text(ws, i, 2, name)
        text(ws, i, 3, content)
    cursor = 7 + len(docs)
    for source in sources["files"]:
        for col, value in enumerate((Path(source["logical_path"]).name, source["logical_path"], source["sha256"], source["bytes"]), 2):
            text(ws, cursor, col, value)
        cursor += 1
    ws.column_dimensions["C"].width = 110
    ws.column_dimensions["D"].width = 70
    finish(wb, ws, cursor - 1, 5)

    ws = sheet(wb, "封面", "206题 · 看清模型在答什么", "V3 C5 seed13 | 正向逐root评分 | AI代理参考标签 | 不是人类金标，不是线上真人准确率", ["指标", "数值", "解释"])
    metrics = [("题目数", '=COUNTA(\'逐题原表\'!$C$7:$C$212)', "同43簇内206个root，反向不加分母", "0"),
               ("与AI参考一致", '=SUM(\'逐题原表\'!$I$7:$I$212)', "合法作答≠正确；对参考一致才计入", "0"),
               ("准确率", "=C8/C7", "所有类别按样本数量计权", "0.0%"),
               ("宏平均F1", '=AVERAGE(\'类别与混淆\'!$H$7:$H$11)', "五个类别同等重要，缺席正确类记0", "0.0000"),
               ("多数类常数准确率", '=MAX(\'类别与混淆\'!$C$7:$C$11)/C7', "事后永远猜相交但互不包含", "0.0%"),
               ("多数类常数宏F1", '=2*MAX(\'类别与混淆\'!$C$7:$C$11)/(C7+MAX(\'类别与混淆\'!$C$7:$C$11))/5', "其余四类召回为零", "0.0000"),
               ("完整输入题数", '=COUNTIF(\'逐题原表\'!$K$7:$K$212,"完整")', "完整只表示模型看全，不保证标答可靠", "0"),
               ("0.80分数通过题数", '=COUNTIF(\'逐题原表\'!$J$7:$J$212,">=0.8")', "本批无通过题；不是零风险认证", "0"),
               ("交换一致题数", '=SUM(\'逐题原表\'!$AA$7:$AA$212)', "正反向按语义交换映射一致的数量", "0")]
    for i, (name, expr, note, fmt) in enumerate(metrics, 7):
        text(ws, i, 2, name)
        formula(ws, i, 3, expr, fmt)
        text(ws, i, 4, note)
    text(ws, 18, 2, "如何阅读")
    for i, name in enumerate([name for name in wb.sheetnames if name != "封面"], 20):
        cell = text(ws, i, 2, name)
        cell.hyperlink = f"#'{name}'!B2"
        text(ws, i, 3, "点击进入；全文阅读可使用同目录HTML")
    text(ws, 30, 2, "诊断计划")
    text(ws, 31, 2, "先核参考可靠性与任务代表性，再测原401训练题是否拟合；之后同样本/同预算单seed控变量，不在这206已曝光题追分。")
    text(ws, 33, 2, "文件不含PivotTable，无需刷新透视；公式请求打开即重算。全部图表为原生Excel对象，图表覆盖原始表的分类与分层汇总。")
    for col, width in (("B", 31), ("C", 20), ("D", 85)):
        ws.column_dimensions[col].width = width
    wb._sheets.insert(0, wb._sheets.pop(wb._sheets.index(ws)))
    finish(wb, ws, 33, 4)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(PATH)
    check = load_workbook(PATH)
    assert check.sheetnames[0] == "封面"
    assert sum(len(sheet._charts) for sheet in check.worksheets) >= 10
    assert check["逐题原表"].max_row == 212
    for i, row in enumerate(rows, 7):
        assert check["逐题原表"].cell(i, 14).value == row["old_text"]
        assert check["逐题原表"].cell(i, 15).value == row["new_text"]
        assert check["逐题原表"].cell(i, 16).value == row["background"]
    check.close()
    exported = [{key: value for key, value in row.items() if key not in ("input_audit", "evidence", "teacher_evidence")} for row in rows]
    payload = json.dumps({"rows": exported, "summary": summary, "classes": CLASSES, "zh": ZH, "families": FAMILY}, ensure_ascii=False).replace("<", "\\u003c")
    template = (ROOT / "reader_template_20261024.html").read_text(encoding="utf-8")
    (ROOT / "206题可视化与逐题阅读器-2026-10-24.html").write_text(template.replace("__PAYLOAD__", payload), encoding="utf-8", newline="\n")
    print(json.dumps({"xlsx": PATH.name, "sheets": wb.sheetnames, "charts": sum(len(sheet._charts) for sheet in wb.worksheets), "raw_rows": len(rows), "fulltext_rows": wb["全文与可见输入"].max_row - 6, "html": "206题可视化与逐题阅读器-2026-10-24.html"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
